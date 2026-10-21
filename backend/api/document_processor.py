import os
import re
import json
import shutil
import subprocess
from datetime import datetime
from PIL import Image, ImageEnhance, ImageFilter
import pymupdf
from docx import Document as DocxDocument
from django.conf import settings
from .models import SourceDocument, ExtractedPage, LogbookEntry
from .ai_provider import is_ai_available, generate_vision_completion, AIProviderUnavailableError


def run_multimodal_vision_analysis(image_path: str) -> tuple[str, float, dict]:
    """
    Analyzes handwritten and scanned logbook pages using Google Gemini Flash Multimodal Vision.
    Extracts transcribed text, week number, daily entries, tools used, and technical observations.
    Returns (transcribed_text, confidence_score, structured_dict).
    Falls back to ("", 0.0, {}) if AI is unavailable or call fails.
    """
    if not is_ai_available() or not os.path.exists(image_path):
        return "", 0.0, {}

    try:
        with open(image_path, "rb") as f:
            image_bytes = f.read()

        ext = os.path.splitext(image_path)[1].lower()
        mime_map = {
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".webp": "image/webp",
        }
        mime_type = mime_map.get(ext, "image/jpeg")

        prompt = (
            "You are an expert academic OCR and SIWES technical logbook analyzer. "
            "Examine this handwritten SIWES industrial training logbook page. "
            "Transcribe all handwritten notes, table entries, and dates accurately.\n\n"
            "Respond ONLY with a JSON object following this exact schema:\n"
            "{\n"
            '  "transcribed_text": "Full verbatim transcription of all text on this page",\n'
            '  "week_number": 1,\n'
            '  "dates_covered": "e.g. June 10 - June 14, 2024",\n'
            '  "entries": [\n'
            '    {\n'
            '      "day_of_week": "Monday",\n'
            '      "activity_summary": "Detailed technical summary of tasks performed, circuits wired, code written, or machines operated",\n'
            '      "hours_spent": 8\n'
            '    }\n'
            '  ],\n'
            '  "tools_and_equipment": ["Oscilloscope", "CAT6 Cable", "Router", "VS Code"],\n'
            '  "problems_encountered": ["Short circuit during bench test"],\n'
            '  "skills_acquired": ["Fiber optic termination"],\n'
            '  "diagram_detected": false,\n'
            '  "supervisor_signed": true,\n'
            '  "confidence": 0.94\n'
            "}"
        )

        response_text = generate_vision_completion(
            prompt=prompt,
            image_bytes=image_bytes,
            mime_type=mime_type,
            response_json=True
        )

        clean_json = response_text.strip()
        if clean_json.startswith("```json"):
            clean_json = clean_json[7:]
        if clean_json.startswith("```"):
            clean_json = clean_json[3:]
        if clean_json.endswith("```"):
            clean_json = clean_json[:-3]
        clean_json = clean_json.strip()

        data = json.loads(clean_json)
        transcribed_text = data.get("transcribed_text", "").strip()
        confidence = float(data.get("confidence", 0.92))
        return transcribed_text, confidence, data
    except Exception:
        return "", 0.0, {}


def run_ocr_on_image(image_path: str) -> tuple[str, float]:
    """
    Attempts to run OCR on an image file using tesseract CLI if installed.
    Returns (extracted_text, confidence_score).
    If OCR is not available, returns empty string and 0.0 confidence.
    """
    tesseract_bin = getattr(settings, 'TESSERACT_CMD', 'tesseract')
    resolved_tesseract = (tesseract_bin if os.path.isabs(tesseract_bin) and os.path.exists(tesseract_bin) else None) or shutil.which(tesseract_bin) or shutil.which('tesseract')
    if not resolved_tesseract:
        return "", 0.0
    tesseract_bin = resolved_tesseract

    try:
        # Preprocess image for OCR using Pillow
        with Image.open(image_path) as img:
            # Convert to grayscale
            gray = img.convert('L')
            # Enhance contrast
            enhancer = ImageEnhance.Contrast(gray)
            enhanced = enhancer.enhance(1.8)
            # Filter noise
            filtered = enhanced.filter(ImageFilter.MedianFilter())
            
            temp_preprocessed = f"{image_path}_preprocessed.png"
            filtered.save(temp_preprocessed)

        # Run tesseract CLI
        result = subprocess.run(
            [tesseract_bin, temp_preprocessed, "stdout", "--oem", "1", "-l", "eng"],
            capture_output=True,
            text=True,
            timeout=30
        )
        
        if os.path.exists(temp_preprocessed):
            try:
                os.remove(temp_preprocessed)
            except OSError:
                pass

        if result.returncode == 0:
            text = result.stdout.strip()
            confidence = 0.82 if len(text) > 30 else 0.50
            return text, confidence
        return "", 0.0
    except Exception:
        return "", 0.0


def extract_logbook_entries_from_text(project, page_instance, text: str):
    """
    Heuristically parses daily or weekly entries from extracted text
    and creates LogbookEntry records for student review.
    """
    if not text or len(text.strip()) < 10:
        return

    lines = [line.strip() for line in text.split('\n') if line.strip()]
    current_week = 1
    
    # Check if a week header exists
    week_match = re.search(r'week\s*[:#-]?\s*(\d+)', text, re.IGNORECASE)
    if week_match:
        try:
            current_week = int(week_match.group(1))
        except ValueError:
            current_week = 1

    # Look for days of week
    day_patterns = [
        ('Monday', r'\b(?:mon|monday)\b'),
        ('Tuesday', r'\b(?:tue|tues|tuesday)\b'),
        ('Wednesday', r'\b(?:wed|wednesday)\b'),
        ('Thursday', r'\b(?:thu|thur|thurs|thursday)\b'),
        ('Friday', r'\b(?:fri|friday)\b'),
        ('Saturday', r'\b(?:sat|saturday)\b'),
    ]

    day_indices = []
    for day_name, pattern in day_patterns:
        for m in re.finditer(pattern, text, re.IGNORECASE):
            day_indices.append((m.start(), day_name))

    day_indices.sort(key=lambda x: x[0])

    if day_indices:
        for i, (start_idx, day_name) in enumerate(day_indices):
            end_idx = day_indices[i + 1][0] if i + 1 < len(day_indices) else len(text)
            chunk = text[start_idx:end_idx].strip()
            # Remove day name itself from summary if at beginning
            clean_chunk = re.sub(r'^(?:mon|monday|tue|tues|tuesday|wed|wednesday|thu|thur|thurs|thursday|fri|friday|sat|saturday)[\s:\-–—]+', '', chunk, flags=re.IGNORECASE).strip()
            summary = clean_chunk[:80].replace('\n', ' ') if clean_chunk else f"{day_name} Activities"
            
            LogbookEntry.objects.create(
                project=project,
                source_page=page_instance,
                week_number=current_week,
                day_of_week=day_name,
                raw_text=clean_chunk or chunk,
                verified_text=clean_chunk or chunk,
                activity_summary=summary,
                is_verified=False,
                order=i
            )
    else:
        # If no explicit days, create an entry for the page content
        summary = lines[0][:80] if lines else f"Page {page_instance.page_number} Activity"
        LogbookEntry.objects.create(
            project=project,
            source_page=page_instance,
            week_number=current_week,
            day_of_week="",
            raw_text=text,
            verified_text=text,
            activity_summary=summary,
            is_verified=False,
            order=LogbookEntry.objects.filter(project=project, week_number=current_week).count()
        )


def process_source_document(doc_id: int):
    """
    Main background or synchronous processing function for an uploaded document.
    """
    try:
        doc = SourceDocument.objects.get(id=doc_id)
    except SourceDocument.DoesNotExist:
        return False

    doc.status = 'processing'
    doc.save(update_fields=['status'])

    file_path = doc.file.path if doc.file else None
    
    try:
        if doc.file_type == 'pdf' and file_path and os.path.exists(file_path):
            pdf = pymupdf.open(file_path)
            doc.page_count = len(pdf)
            doc.save(update_fields=['page_count'])

            for page_index in range(len(pdf)):
                page_num = page_index + 1
                page = pdf[page_index]
                text = page.get_text("text").strip()
                confidence = 0.95
                page_status = 'processed'
                notes = ''

                # If text is too short, page might be scanned or image-based
                if len(text) < 30:
                    pix = page.get_pixmap(dpi=150)
                    temp_img_path = f"{file_path}_p{page_num}.png"
                    pix.save(temp_img_path)
                    
                    v_text, v_conf, v_data = run_multimodal_vision_analysis(temp_img_path)
                    if v_text:
                        text = v_text
                        confidence = v_conf
                        page_status = 'processed'
                        notes = 'Analyzed with Gemini Flash Multimodal Vision'
                    else:
                        ocr_text, ocr_conf = run_ocr_on_image(temp_img_path)
                        if ocr_text:
                            text = ocr_text
                            confidence = ocr_conf
                            page_status = 'processed'
                        else:
                            confidence = 0.20
                            page_status = 'manual_correction_needed'
                            notes = 'Page appears to be scanned or handwritten. Please review and input or verify the text.'

                    if os.path.exists(temp_img_path):
                        try:
                            os.remove(temp_img_path)
                        except OSError:
                            pass

                extracted_page = ExtractedPage.objects.create(
                    document=doc,
                    page_number=page_num,
                    extracted_text=text,
                    confidence_score=confidence,
                    status=page_status,
                    notes=notes,
                    is_reviewed=False
                )

                if text:
                    extract_logbook_entries_from_text(doc.project, extracted_page, text)

            pdf.close()

        elif doc.file_type == 'image' and file_path and os.path.exists(file_path):
            doc.page_count = 1
            doc.save(update_fields=['page_count'])
            
            vision_text, vision_conf, vision_data = run_multimodal_vision_analysis(file_path)
            if vision_text:
                page_status = 'processed'
                confidence = vision_conf
                text = vision_text
                notes = 'Analyzed with Gemini Flash Multimodal Vision'
            else:
                ocr_text, ocr_conf = run_ocr_on_image(file_path)
                if ocr_text:
                    page_status = 'processed'
                    confidence = ocr_conf
                    text = ocr_text
                    notes = ''
                else:
                    page_status = 'manual_correction_needed'
                    confidence = 0.25
                    text = ''
                    notes = 'Scanned / handwritten image uploaded. Please verify and transcribe handwritten text.'

            extracted_page = ExtractedPage.objects.create(
                document=doc,
                page_number=1,
                extracted_text=text,
                confidence_score=confidence,
                status=page_status,
                notes=notes,
                is_reviewed=False
            )
            
            # If Vision extracted structured daily entries, create them with high confidence
            if vision_data and vision_data.get('entries'):
                week_num = vision_data.get('week_number') or 1
                for idx, entry in enumerate(vision_data.get('entries', [])):
                    summary = entry.get('activity_summary', '').strip()
                    day = entry.get('day_of_week', f'Day {idx+1}')
                    if summary:
                        LogbookEntry.objects.create(
                            project=doc.project,
                            source_page=extracted_page,
                            week_number=week_num,
                            day_of_week=day,
                            raw_text=summary,
                            verified_text=summary,
                            activity_summary=summary[:120],
                            order=idx + 1,
                            is_verified=True
                        )
                # Auto-enrich project tools if detected
                tools = vision_data.get('tools_and_equipment', [])
                if tools and not doc.project.tools_used:
                    doc.project.tools_used = ', '.join(tools)
                    doc.project.save(update_fields=['tools_used'])
            elif text:
                extract_logbook_entries_from_text(doc.project, extracted_page, text)

        elif doc.file_type == 'docx' and file_path and os.path.exists(file_path):
            doc_obj = DocxDocument(file_path)
            paragraphs = [p.text.strip() for p in doc_obj.paragraphs if p.text.strip()]
            full_text = '\n\n'.join(paragraphs)
            
            # Also extract text from any tables
            table_texts = []
            for table in doc_obj.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        table_texts.append(row_text)
            
            if table_texts:
                full_text += '\n\nTable Data:\n' + '\n'.join(table_texts)

            doc.page_count = 1
            doc.save(update_fields=['page_count'])

            extracted_page = ExtractedPage.objects.create(
                document=doc,
                page_number=1,
                extracted_text=full_text,
                confidence_score=0.98,
                status='processed',
                notes='Direct Word extraction',
                is_reviewed=False
            )
            if full_text:
                extract_logbook_entries_from_text(doc.project, extracted_page, full_text)

        elif doc.file_type == 'note':
            # Note with direct text provided
            page = ExtractedPage.objects.create(
                document=doc,
                page_number=1,
                extracted_text=doc.error_message or "", # Note content stored temporarily in error_message or separate note
                confidence_score=1.0,
                status='processed',
                notes='Typed note entered directly',
                is_reviewed=True
            )
            if doc.error_message:
                extract_logbook_entries_from_text(doc.project, page, doc.error_message)
            doc.error_message = ""

        doc.status = 'completed'
        doc.save(update_fields=['status', 'error_message'])
        
        # Advance project status if still draft/uploading
        if doc.project.status in ['draft', 'uploading']:
            doc.project.status = 'reviewing_sources'
            doc.project.save(update_fields=['status'])
            
        return True

    except Exception as e:
        doc.status = 'error'
        doc.error_message = str(e)
        doc.save(update_fields=['status', 'error_message'])
        return False
