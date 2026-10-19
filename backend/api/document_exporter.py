import os
import re
import shutil
import subprocess
from pathlib import Path
from html.parser import HTMLParser
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from django.conf import settings
from django.utils import timezone
from .models import Report, ReportSection


class SimpleHTMLDocxParser(HTMLParser):
    """
    Parses basic HTML tags (<p>, <h3>, <h4>, <ul>, <ol>, <li>, <strong>, <em>, <table>, <tr>, <th>, <td>)
    and renders them directly into a python-docx Document with proper academic formatting.
    """
    def __init__(self, doc, default_align=WD_ALIGN_PARAGRAPH.LEFT):
        super().__init__()
        self.doc = doc
        self.default_align = default_align
        self.current_p = None
        self.bold = False
        self.italic = False
        self.list_type = None
        self.in_table = False
        self.table_rows = []
        self.current_row_cells = []
        self.current_cell_text = []

    def handle_starttag(self, tag, attrs):
        attrs_dict = dict(attrs)
        align_str = attrs_dict.get('style', '')

        if tag == 'table':
            self.in_table = True
            self.table_rows = []
            self.current_row_cells = []
            self.current_cell_text = []
            return

        if self.in_table:
            if tag == 'tr':
                self.current_row_cells = []
            elif tag in ['th', 'td']:
                self.current_cell_text = []
                if tag == 'th':
                    self.bold = True
            return

        if tag in ['p', 'div']:
            self.current_p = self.doc.add_paragraph()
            self.current_p.paragraph_format.line_spacing = 1.5
            self.current_p.paragraph_format.space_after = Pt(6)
            if 'text-align: center' in align_str:
                self.current_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif 'text-align: right' in align_str:
                self.current_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                self.current_p.alignment = self.default_align

        elif tag in ['h1', 'h2']:
            self.current_p = self.doc.add_paragraph()
            self.current_p.paragraph_format.line_spacing = 1.5
            self.current_p.paragraph_format.space_before = Pt(12)
            self.current_p.paragraph_format.space_after = Pt(6)
            self.current_p.alignment = WD_ALIGN_PARAGRAPH.CENTER if 'text-align: center' in align_str else WD_ALIGN_PARAGRAPH.LEFT
            self.bold = True

        elif tag in ['h3', 'h4']:
            self.current_p = self.doc.add_paragraph()
            self.current_p.paragraph_format.line_spacing = 1.5
            self.current_p.paragraph_format.space_before = Pt(8)
            self.current_p.paragraph_format.space_after = Pt(4)
            self.bold = True

        elif tag in ['strong', 'b']:
            self.bold = True

        elif tag in ['em', 'i']:
            self.italic = True

        elif tag == 'ul':
            self.list_type = 'List Bullet'

        elif tag == 'ol':
            self.list_type = 'List Number'

        elif tag == 'li':
            self.current_p = self.doc.add_paragraph(style=self.list_type or 'List Bullet')
            self.current_p.paragraph_format.line_spacing = 1.5
            self.current_p.paragraph_format.space_after = Pt(4)

        elif tag == 'br':
            if self.current_p:
                self.current_p.add_run().add_break()

    def handle_endtag(self, tag):
        if self.in_table:
            if tag in ['th', 'td']:
                cell_str = " ".join(self.current_cell_text).strip()
                self.current_row_cells.append(cell_str)
                self.current_cell_text = []
                self.bold = False
            elif tag == 'tr':
                if self.current_row_cells:
                    self.table_rows.append(self.current_row_cells)
                    self.current_row_cells = []
            elif tag == 'table':
                self.in_table = False
                if self.table_rows:
                    num_cols = max(len(r) for r in self.table_rows)
                    num_rows = len(self.table_rows)
                    if num_cols > 0 and num_rows > 0:
                        tbl = self.doc.add_table(rows=num_rows, cols=num_cols)
                        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                        tblPr = tbl._tbl.tblPr
                        borders = parse_xml(
                            f'<w:tblBorders {nsdecls("w")}>'
                            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                            f'<w:left w:val="none"/>'
                            f'<w:right w:val="none"/>'
                            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                            f'<w:insideV w:val="none"/>'
                            f'</w:tblBorders>'
                        )
                        tblPr.append(borders)
                        for r_idx, row_data in enumerate(self.table_rows):
                            for c_idx, cell_data in enumerate(row_data):
                                if c_idx < num_cols:
                                    cell = tbl.cell(r_idx, c_idx)
                                    p = cell.paragraphs[0]
                                    p.paragraph_format.line_spacing = 1.15
                                    p.paragraph_format.space_after = Pt(2)
                                    run = p.add_run(cell_data)
                                    run.font.name = 'Times New Roman'
                                    run.font.size = Pt(10)
                                    if r_idx == 0:
                                        run.bold = True
                        p_spacer = self.doc.add_paragraph()
                        p_spacer.paragraph_format.space_before = Pt(4)
                        p_spacer.paragraph_format.space_after = Pt(4)
                self.table_rows = []
                self.current_row_cells = []
                self.current_cell_text = []
            return

        if tag in ['strong', 'b']:
            self.bold = False
        elif tag in ['em', 'i']:
            self.italic = False
        elif tag in ['ul', 'ol']:
            self.list_type = None

    def handle_data(self, data):
        cleaned = data.strip()
        if not cleaned:
            return

        if self.in_table:
            self.current_cell_text.append(data.strip())
            return

        if self.current_p is None:
            self.current_p = self.doc.add_paragraph()
            self.current_p.paragraph_format.line_spacing = 1.5
            self.current_p.paragraph_format.space_after = Pt(6)
            self.current_p.alignment = self.default_align

        run = self.current_p.add_run(data)
        run.bold = self.bold
        run.italic = self.italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)


def set_academic_margins(section):
    """
    Applies standard Nigerian academic margins for binding:
    Left: 1.5 inches (binding side)
    Right: 1.0 inch
    Top: 1.0 inch
    Bottom: 1.0 inch
    Page size: A4 (8.27 x 11.69 inches)
    """
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)


def generate_docx_report(report: Report) -> str:
    """
    Builds a professional, properly formatted A4 Word document for the SIWES report.
    Returns the absolute path to the generated DOCX file.
    """
    doc = Document()
    
    # Configure base style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    # First section: preliminary pages
    first_section = doc.sections[0]
    set_academic_margins(first_section)

    sections = report.sections.all().order_by('chapter_number', 'order', 'id')
    current_chapter = None

    for idx, sec in enumerate(sections):
        # Chapter 0 preliminary pages get page break between each
        # Chapters 1..5 get page break before Chapter starts
        if sec.chapter_number == 0:
            if idx > 0:
                doc.add_page_break()
        else:
            if current_chapter != sec.chapter_number:
                doc.add_page_break()
                current_chapter = sec.chapter_number
                
                ch_titles = {
                    1: "CHAPTER ONE\nINTRODUCTION",
                    2: "CHAPTER TWO\nORGANIZATION PROFILE",
                    3: "CHAPTER THREE\nTRAINING ACTIVITIES AND TECHNICAL EXPERIENCE",
                    4: "CHAPTER FOUR\nSKILLS ACQUIRED, CHALLENGES AND SOLUTIONS",
                    5: "CHAPTER FIVE\nCONCLUSION AND RECOMMENDATIONS",
                    6: "REFERENCES AND APPENDICES"
                }
                header_text = ch_titles.get(sec.chapter_number, f"CHAPTER {sec.chapter_number}")
                ch_p = doc.add_paragraph()
                ch_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ch_p.paragraph_format.space_before = Pt(18)
                ch_p.paragraph_format.space_after = Pt(18)
                run = ch_p.add_run(header_text)
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(14)

        if sec.chapter_number > 0 and sec.section_key not in ["references", "appendices"]:
            sub_p = doc.add_paragraph()
            sub_p.paragraph_format.space_before = Pt(12)
            sub_p.paragraph_format.space_after = Pt(6)
            sub_run = sub_p.add_run(sec.title)
            sub_run.bold = True
            sub_run.font.name = 'Times New Roman'
            sub_run.font.size = Pt(12)

        content = sec.content or ""
        parser = SimpleHTMLDocxParser(doc)
        parser.feed(content)

    export_dir = Path(settings.MEDIA_ROOT) / "projects" / str(report.project.id) / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    
    clean_title = re.sub(r'[^a-zA-Z0-9_\-]', '_', report.project.title or "SIWES_Report")
    docx_filename = f"{clean_title}_v{report.version}.docx"
    docx_path = export_dir / docx_filename
    
    doc.save(str(docx_path))
    
    rel_path = f"projects/{report.project.id}/exports/{docx_filename}"
    report.docx_file.name = rel_path
    report.last_exported_at = timezone.now()
    report.save(update_fields=['docx_file', 'last_exported_at'])

    return str(docx_path)


def convert_docx_to_pdf(docx_path: str) -> str:
    """
    Renders DOCX into PDF using headless LibreOffice.
    Validates output PDF size and magic bytes.
    Returns the absolute path to the PDF.
    """
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"Source DOCX file not found: {docx_path}")

    out_dir = str(Path(docx_path).parent)
    libreoffice_bin = getattr(settings, 'LIBREOFFICE_PATH', 'libreoffice')

    cmd = [
        libreoffice_bin,
        '--headless',
        '--convert-to', 'pdf',
        '--outdir', out_dir,
        docx_path
    ]

    result = subprocess.run(cmd, capture_output=True, text=True, timeout=90)
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice PDF conversion failed: {result.stderr or result.stdout}")

    pdf_path = os.path.splitext(docx_path)[0] + '.pdf'
    
    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"PDF was not generated at expected path: {pdf_path}")

    file_size = os.path.getsize(pdf_path)
    if file_size < 1000:
        raise ValueError(f"Generated PDF is suspiciously small ({file_size} bytes). Conversion may be corrupted.")

    with open(pdf_path, 'rb') as f:
        header = f.read(5)
        if not header.startswith(b'%PDF-'):
            raise ValueError("Generated file is not a valid PDF document (missing %PDF header).")

    return pdf_path


def export_report_documents(report: Report) -> dict:
    """
    Executes both DOCX generation and PDF conversion with full validation.
    Returns dict with file paths and metadata.
    """
    docx_path = generate_docx_report(report)
    pdf_path = convert_docx_to_pdf(docx_path)

    pdf_filename = os.path.basename(pdf_path)
    rel_pdf_path = f"projects/{report.project.id}/exports/{pdf_filename}"
    
    report.pdf_file.name = rel_pdf_path
    report.project.status = 'exported'
    report.project.save(update_fields=['status'])
    report.save(update_fields=['pdf_file'])

    return {
        "success": True,
        "docx_path": docx_path,
        "pdf_path": pdf_path,
        "docx_filename": os.path.basename(docx_path),
        "pdf_filename": pdf_filename,
        "size_bytes": os.path.getsize(pdf_path),
        "exported_at": report.last_exported_at.isoformat() if report.last_exported_at else None
    }
